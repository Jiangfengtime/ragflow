#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
import re
import os
from functools import reduce
from io import BytesIO
from timeit import default_timer as timer
from typing import Any, Callable
from docx import Document
from docx.opc.pkgreader import _SerializedRelationships, _SerializedRelationship
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from docx.opc.oxml import parse_xml
from markdown import markdown
from PIL import Image
from common.token_utils import num_tokens_from_string

from common.constants import LLMType, MAXIMUM_PAGE_NUMBER
from api.db.services.llm_service import LLMBundle
from api.db.joint_services.tenant_model_service import (
    ensure_mineru_from_env,
    ensure_opendataloader_from_env,
    ensure_paddleocr_from_env,
    get_composite_model_name_by_id,
    get_first_provider_model_name,
    resolve_model_config,
    get_tenant_default_model_by_type,
)
from rag.utils.file_utils import extract_embed_file, extract_links_from_pdf, extract_links_from_docx, extract_html
from deepdoc.parser import DocxParser, EpubParser, ExcelParser, HtmlParser, JsonParser, MarkdownElementExtractor, MarkdownParser, PdfParser, TxtParser
from deepdoc.parser.figure_parser import VisionFigureParser, vision_figure_parser_docx_wrapper_naive, vision_figure_parser_pdf_wrapper
from deepdoc.parser.pdf_parser import PlainParser, VisionParser
from deepdoc.parser.docling_parser import DoclingParser
from deepdoc.parser.tcadp_parser import TCADPParser
from common.float_utils import normalize_overlapped_percent
from common.parser_config_utils import has_mineru_options, normalize_layout_recognizer
from common.text_utils import normalize_arabic_presentation_forms
from rag.nlp import (
    concat_img,
    find_codec,
    naive_merge,
    naive_merge_with_images,
    naive_merge_docx,
    rag_tokenizer,
    tokenize_chunks,
    doc_tokenize_chunks_with_images,
    tokenize_table,
    append_context2table_image4pdf,
    tokenize_chunks_with_images,
)  # noqa: F401


def _is_short_header(text, max_tokens=50):
    """
    Check if text is a short markdown header.

    Args:
        text: The text to check
        max_tokens: Maximum tokens for a header to be considered "short"

    Returns:
        bool: True if text is a short markdown header, False otherwise
    """
    if not text or not text.strip():
        return False

    # Check if it matches markdown header pattern: 1-6 # followed by space
    if not re.match(r"^#{1,6}\s+", text.strip()):
        return False

    # Check if token count is below threshold
    return num_tokens_from_string(text) < max_tokens


def _normalize_section_text_for_rtl_presentation_forms(sections):
    if not sections:
        return sections

    normalized_sections = []
    for section in sections:
        if isinstance(section, tuple):
            if not section:
                normalized_sections.append(section)
                continue
            text = section[0]
            normalized_text = normalize_arabic_presentation_forms(text)
            normalized_sections.append((normalized_text, *section[1:]))
            continue
        if isinstance(section, list):
            if not section:
                normalized_sections.append(section)
                continue
            text = section[0]
            normalized_text = normalize_arabic_presentation_forms(text)
            normalized_sections.append([normalized_text, *section[1:]])
            continue
        normalized_sections.append(normalize_arabic_presentation_forms(section))

    return normalized_sections


def by_deepdoc(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    pdf_parser = pdf_cls() if pdf_cls else Pdf()
    sections, tables = pdf_parser(filename if binary is None else binary, from_page=from_page, to_page=to_page, callback=callback)

    tables = vision_figure_parser_pdf_wrapper(
        tbls=tables,
        sections=sections,
        callback=callback,
        lang=lang,
        **kwargs,
    )
    return sections, tables, pdf_parser


def _dispatch_pdf_parser(parser_config: dict, opendataloader_llm_name=None, layout_recognize_override: str | None = None) -> tuple[Callable, str, Any, str, Any]:
    """Resolve the PDF parser callable for the current ``parser_config``.

    Returns a 5-tuple ``(parser_callable, parser_name, layout_recognizer,
    opendataloader_llm_name, parser_model_name)`` so the dispatch logic
    stays testable in isolation (issue #17114).

    ``layout_recognize_override`` lets callers pass an already-resolved
    layout_recognize value (e.g. a UUID resolved via
    :func:`get_composite_model_name_by_id`) so the dispatch doesn't
    re-read the stale UUID from ``parser_config`` and re-normalize it.

    When ``layout_recognize`` is a value that does not match any known parser
    name — typically a stale ``TenantModel`` UUID stored on the document —
    and the parser config carries MinerU-specific options (``mineru_*``
    keys), the dispatch falls back to :func:`by_mineru` instead of routing
    to :func:`by_plaintext`, which would otherwise try to resolve the UUID
    as an IMAGE2TEXT vision model and crash.
    """
    raw_layout_recognize = layout_recognize_override if layout_recognize_override is not None else parser_config.get("layout_recognize", "DeepDOC")
    layout_recognizer, parser_model_name = normalize_layout_recognizer(raw_layout_recognize)
    if layout_recognizer == "OpenDataLoader" and parser_model_name:
        opendataloader_llm_name = parser_model_name

    if isinstance(layout_recognizer, bool):
        layout_recognizer = "DeepDOC" if layout_recognizer else "PlainText"

    # Normalize "Plain Text" to "plaintext" so the PARSERS lookup below
    # hits the explicit "plaintext" entry instead of falling through to
    # the by_plaintext default — important because the MinerU fallback
    # guard below keys off whether the parsed name is a known keyword.
    if layout_recognizer == "Plain Text":
        layout_recognizer = "plaintext"
    name = layout_recognizer.strip().lower()
    parser = PARSERS.get(name, by_plaintext)

    # Closes #17114: when the document's layout_recognize is a model id
    # (e.g. a TenantModel UUID) that does not match any known parser name,
    # the previous dispatch fell through to by_plaintext, which tried to
    # resolve the id as an IMAGE2TEXT vision model and failed with
    # ``Provider <empty> not found for model <id>``. If mineru-specific
    # options are set in parser_config, the operator clearly intended the
    # MinerU parser, so route there instead and surface a clear log line
    # rather than masking the misconfiguration silently.
    # Guard: only fall back when the parser name is NOT a known keyword
    # (e.g. "DeepDOC", "Plain Text"). A configuration like
    # ``{"layout_recognize": "Plain Text", "mineru_lang": "English"}``
    # must keep honoring PlainText, not be silently rerouted to MinerU.
    if name not in PARSERS and parser is by_plaintext and has_mineru_options(parser_config):
        logging.warning(
            "[naive] layout_recognize=%r does not match a known parser; falling back to MinerU because mineru_* options are set (see issue #17114).",
            layout_recognizer,
        )
        parser = by_mineru
        name = "mineru"

    return parser, name, layout_recognizer, opendataloader_llm_name, parser_model_name


def by_mineru(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    mineru_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not mineru_llm_name:
            try:
                mineru_llm_name = get_first_provider_model_name(tenant_id, "MinerU", LLMType.OCR) or ensure_mineru_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env mineru: {e}")

        if mineru_llm_name:
            try:
                ocr_model_config = resolve_model_config(tenant_id, LLMType.OCR, mineru_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl

                # Closes #14869: when the tenant has a VISION model
                # configured, let the MinerU parser enrich image chunks with
                # VLM-generated semantic descriptions (parity with deepdoc's
                # VisionFigureParser). Best-effort — fall back silently if
                # no vision model is available.
                if "vision_model" not in kwargs:
                    try:
                        vision_model_config = get_tenant_default_model_by_type(tenant_id, LLMType.VISION)
                        kwargs["vision_model"] = LLMBundle(tenant_id=tenant_id, model_config=vision_model_config, lang=lang)
                    except Exception as vlm_err:
                        logging.info(f"[MinerU] no VISION model for tenant; skipping image VLM enhancement: {vlm_err}")

                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    lang=lang,
                    page_from=from_page,
                    page_to=min(to_page, MAXIMUM_PAGE_NUMBER),
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle MinerU ({mineru_llm_name}): {e}")
                raise

    raise RuntimeError("MinerU model not found or not configured.")


def by_docling(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    pdf_parser = DoclingParser()
    parse_method = kwargs.get("parse_method", "raw")

    if not pdf_parser.check_installation():
        if callback:
            callback(-1, "Docling not found.")
        return None, None, pdf_parser

    sections, tables = pdf_parser.parse_pdf(
        filepath=filename,
        binary=binary,
        callback=callback,
        output_dir=os.environ.get("DOCLING_OUTPUT_DIR", ""),
        delete_output=bool(int(os.environ.get("DOCLING_DELETE_OUTPUT", 1))),
        docling_server_url=os.environ.get("DOCLING_SERVER_URL", ""),
        parse_method=parse_method,
    )
    return sections, tables, pdf_parser


def by_opendataloader(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    opendataloader_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    if tenant_id:
        if not opendataloader_llm_name:
            try:
                opendataloader_llm_name = get_first_provider_model_name(tenant_id, "OpenDataLoader", LLMType.OCR) or ensure_opendataloader_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env opendataloader: {e}")

        if opendataloader_llm_name:
            try:
                ocr_model_config = resolve_model_config(tenant_id, LLMType.OCR, opendataloader_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                parse_options = {k: kwargs[k] for k in ("hybrid", "image_output", "sanitize") if k in kwargs}
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **parse_options,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle OpenDataLoader ({opendataloader_llm_name}): {e}")

    if callback:
        callback(-1, "OpenDataLoader not found.")
    return None, None, None


def by_tcadp(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, pdf_cls=None, **kwargs):
    tcadp_parser = TCADPParser()

    if not tcadp_parser.check_installation():
        callback(-1, "TCADP parser not available. Please check Tencent Cloud API configuration.")
        return None, None, tcadp_parser

    sections, tables = tcadp_parser.parse_pdf(filepath=filename, binary=binary, callback=callback, output_dir=os.environ.get("TCADP_OUTPUT_DIR", ""), file_type="PDF")
    return sections, tables, tcadp_parser


def by_paddleocr(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    paddleocr_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not paddleocr_llm_name:
            try:
                paddleocr_llm_name = get_first_provider_model_name(tenant_id, "PaddleOCR", LLMType.OCR) or ensure_paddleocr_from_env(tenant_id)
            except Exception as e:  # best-effort fallback
                logging.warning(f"fallback to env paddleocr: {e}")

        if paddleocr_llm_name:
            try:
                ocr_model_config = resolve_model_config(tenant_id, LLMType.OCR, paddleocr_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle PaddleOCR ({paddleocr_llm_name}): {e}")

        return None, None, None

    if callback:
        callback(-1, "PaddleOCR not found.")
    return None, None, None


def by_somark(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    somark_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not somark_llm_name:
            try:
                from api.db.joint_services.tenant_model_service import ensure_somark_from_env

                somark_llm_name = ensure_somark_from_env(tenant_id)
            except Exception as e:
                logging.warning(f"fallback to env somark: {e}")

        if somark_llm_name:
            try:
                ocr_model_config = resolve_model_config(tenant_id, LLMType.OCR, somark_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle SoMark ({somark_llm_name}): {e}")
                if callback:
                    callback(-1, f"Failed to parse pdf via SoMark ({somark_llm_name}): {e}")
                return None, None, None

    if callback:
        callback(-1, "SoMark not found.")
    return None, None, None


def by_mistral_ocr(
    filename,
    binary=None,
    from_page=0,
    to_page=MAXIMUM_PAGE_NUMBER,
    lang="Chinese",
    callback=None,
    pdf_cls=None,
    parse_method: str = "raw",
    mistral_ocr_llm_name: str | None = None,
    tenant_id: str | None = None,
    **kwargs,
):
    pdf_parser = None
    if tenant_id:
        if not mistral_ocr_llm_name:
            try:
                from api.db.joint_services.tenant_model_service import ensure_mistral_ocr_from_env

                mistral_ocr_llm_name = ensure_mistral_ocr_from_env(tenant_id)
            except Exception as e:
                logging.warning(f"fallback to env mistral ocr: {e}")

        if mistral_ocr_llm_name:
            try:
                ocr_model_config = resolve_model_config(tenant_id, LLMType.OCR, mistral_ocr_llm_name)
                ocr_model = LLMBundle(tenant_id=tenant_id, model_config=ocr_model_config, lang=lang)
                pdf_parser = ocr_model.mdl
                # Best-effort figure description: hand the parser the tenant's
                # vision model so Mistral OCR's extracted images get VLM captions
                # (parity with MinerU/deepdoc). Skip silently if none is configured.
                if "vision_model" not in kwargs:
                    try:
                        vision_model_config = get_tenant_default_model_by_type(tenant_id, LLMType.VISION)
                        kwargs["vision_model"] = LLMBundle(tenant_id=tenant_id, model_config=vision_model_config, lang=lang)
                    except Exception as vlm_err:
                        logging.info(f"[Mistral OCR] no vision model for tenant; skipping figure description: {vlm_err}")
                sections, tables = pdf_parser.parse_pdf(
                    filepath=filename,
                    binary=binary,
                    callback=callback,
                    parse_method=parse_method,
                    from_page=from_page,
                    to_page=to_page,
                    lang=lang,
                    **kwargs,
                )
                return sections, tables, pdf_parser
            except Exception as e:
                logging.error(f"Failed to parse pdf via LLMBundle Mistral OCR ({mistral_ocr_llm_name}): {e}")
                if callback:
                    callback(-1, f"Failed to parse pdf via Mistral OCR ({mistral_ocr_llm_name}): {e}")
                return None, None, None

    if callback:
        callback(-1, "Mistral OCR not found.")
    return None, None, None


def by_plaintext(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, callback=None, **kwargs):
    layout_recognizer = (kwargs.get("layout_recognizer") or "").strip()
    if (not layout_recognizer) or layout_recognizer.replace(" ", "").lower() == "plaintext":
        pdf_parser = PlainParser()
    else:
        tenant_id = kwargs.get("tenant_id")
        if not tenant_id:
            raise ValueError("tenant_id is required when using vision layout recognizer")
        vision_model_config = resolve_model_config(tenant_id, LLMType.VISION, layout_recognizer)
        vision_model = LLMBundle(
            tenant_id,
            model_config=vision_model_config,
            lang=kwargs.get("lang", "Chinese"),
        )
        pdf_parser = VisionParser(vision_model=vision_model, **kwargs)

    sections, tables = pdf_parser(filename if binary is None else binary, from_page=from_page, to_page=to_page, callback=callback)
    return sections, tables, pdf_parser


PARSERS = {
    "deepdoc": by_deepdoc,
    "mineru": by_mineru,
    "docling": by_docling,
    "opendataloader": by_opendataloader,
    "tcadp parser": by_tcadp,
    "paddleocr": by_paddleocr,
    "somark": by_somark,
    "mistral ocr": by_mistral_ocr,
    "plaintext": by_plaintext,  # default
}


class Docx(DocxParser):
    def __init__(self):
        pass

    def __clean(self, line):
        line = re.sub(r"\u3000", " ", line).strip()
        return line

    def __get_nearest_title(self, table_index, filename):
        """Get the hierarchical title structure before the table"""
        import re
        from docx.text.paragraph import Paragraph

        titles = []
        blocks = []

        # Get document name from filename parameter
        doc_name = re.sub(r"\.[a-zA-Z]+$", "", filename)
        if not doc_name:
            doc_name = "Untitled Document"

        # Collect all document blocks while maintaining document order
        try:
            # Iterate through all paragraphs and tables in document order
            for i, block in enumerate(self.doc._element.body):
                if block.tag.endswith("p"):  # Paragraph
                    p = Paragraph(block, self.doc)
                    blocks.append(("p", i, p))
                elif block.tag.endswith("tbl"):  # Table
                    blocks.append(("t", i, None))  # Table object will be retrieved later
        except Exception as e:
            logging.error(f"Error collecting blocks: {e}")
            return ""

        # Find the target table position
        target_table_pos = -1
        table_count = 0
        for i, (block_type, pos, _) in enumerate(blocks):
            if block_type == "t":
                if table_count == table_index:
                    target_table_pos = pos
                    break
                table_count += 1

        if target_table_pos == -1:
            return ""  # Target table not found

        # Find the nearest heading paragraph in reverse order
        nearest_title = None
        for i in range(len(blocks) - 1, -1, -1):
            block_type, pos, block = blocks[i]
            if pos >= target_table_pos:  # Skip blocks after the table
                continue

            if block_type != "p":
                continue

            if block.style and block.style.name and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                try:
                    level_match = re.search(r"(\d+)", block.style.name)
                    if level_match:
                        level = int(level_match.group(1))
                        if level <= 7:  # Support up to 7 heading levels
                            title_text = block.text.strip()
                            if title_text:  # Avoid empty titles
                                nearest_title = (level, title_text)
                                break
                except Exception as e:
                    logging.error(f"Error parsing heading level: {e}")

        if nearest_title:
            # Add current title
            titles.append(nearest_title)
            current_level = nearest_title[0]

            # Find all parent headings, allowing cross-level search
            while current_level > 1:
                found = False
                for i in range(len(blocks) - 1, -1, -1):
                    block_type, pos, block = blocks[i]
                    if pos >= target_table_pos:  # Skip blocks after the table
                        continue

                    if block_type != "p":
                        continue

                    if block.style and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                        try:
                            level_match = re.search(r"(\d+)", block.style.name)
                            if level_match:
                                level = int(level_match.group(1))
                                # Find any heading with a higher level
                                if level < current_level:
                                    title_text = block.text.strip()
                                    if title_text:  # Avoid empty titles
                                        titles.append((level, title_text))
                                        current_level = level
                                        found = True
                                        break
                        except Exception as e:
                            logging.error(f"Error parsing parent heading: {e}")

                if not found:  # Break if no parent heading is found
                    break

            # Sort by level (ascending, from highest to lowest)
            titles.sort(key=lambda x: x[0])
            # Organize titles (from highest to lowest)
            hierarchy = [doc_name] + [t[1] for t in titles]
            return " > ".join(hierarchy)

        return ""

    def __call__(self, filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER):
        self.doc = Document(filename) if binary is None else Document(BytesIO(binary))
        pn = 0
        lines = []
        last_image = None
        table_idx = 0

        def flush_last_image():
            nonlocal last_image, lines
            if last_image is not None:
                lines.append({"text": "", "image": last_image, "table": None, "style": "Image"})
                last_image = None

        for block in self.doc._element.body:
            if pn > to_page:
                break

            if block.tag.endswith("p"):
                p = Paragraph(block, self.doc)

                if from_page <= pn < to_page:
                    text = p.text.strip()
                    style_name = p.style.name if p.style else ""

                    if text:
                        if style_name == "Caption":
                            former_image = None

                            if lines and lines[-1].get("image") and lines[-1].get("style") != "Caption":
                                former_image = lines[-1].get("image")
                                lines.pop()

                            elif last_image is not None:
                                former_image = last_image
                                last_image = None

                            lines.append(
                                {
                                    "text": self.__clean(text),
                                    "image": former_image if former_image else None,
                                    "table": None,
                                }
                            )

                        else:
                            flush_last_image()
                            lines.append(
                                {
                                    "text": self.__clean(text),
                                    "image": None,
                                    "table": None,
                                }
                            )

                            current_image = self.get_picture(self.doc, p)
                            if current_image is not None:
                                lines.append(
                                    {
                                        "text": "",
                                        "image": current_image,
                                        "table": None,
                                    }
                                )

                    else:
                        current_image = self.get_picture(self.doc, p)
                        if current_image is not None:
                            last_image = current_image

                for run in p.runs:
                    xml = run._element.xml
                    if "lastRenderedPageBreak" in xml:
                        pn += 1
                        continue
                    if "w:br" in xml and 'type="page"' in xml:
                        pn += 1

            elif block.tag.endswith("tbl"):
                if pn < from_page or pn > to_page:
                    table_idx += 1
                    continue

                flush_last_image()
                tb = DocxTable(block, self.doc)
                title = self.__get_nearest_title(table_idx, filename)
                html = "<table>"
                if title:
                    html += f"<caption>Table Location: {title}</caption>"
                for r in tb.rows:
                    html += "<tr>"
                    col_idx = 0
                    try:
                        while col_idx < len(r.cells):
                            span = 1
                            c = r.cells[col_idx]
                            for j in range(col_idx + 1, len(r.cells)):
                                if c.text == r.cells[j].text:
                                    span += 1
                                    col_idx = j
                                else:
                                    break
                            col_idx += 1
                            html += f"<td>{c.text}</td>" if span == 1 else f"<td colspan='{span}'>{c.text}</td>"
                    except Exception as e:
                        logging.warning(f"Error parsing table, ignore: {e}")
                    html += "</tr>"
                html += "</table>"
                lines.append({"text": "", "image": None, "table": html})
                table_idx += 1

        flush_last_image()
        new_line = [(line.get("text"), line.get("image"), line.get("table")) for line in lines]

        return new_line

    def to_markdown(self, filename=None, binary=None, inline_images: bool = True):
        """
        This function uses mammoth, licensed under the BSD 2-Clause License.
        """

        import base64
        import uuid

        import mammoth
        from markdownify import markdownify

        docx_file = BytesIO(binary) if binary is not None else open(filename, "rb")

        def _convert_image_to_base64(image):
            try:
                with image.open() as image_file:
                    image_bytes = image_file.read()
                encoded = base64.b64encode(image_bytes).decode("utf-8")
                base64_url = f"data:{image.content_type};base64,{encoded}"

                alt_name = "image"
                alt_name = f"img_{uuid.uuid4().hex[:8]}"

                return {"src": base64_url, "alt": alt_name}
            except Exception as e:
                logging.warning(f"Failed to convert image to base64: {e}")
                return {"src": "", "alt": "image"}

        try:
            if inline_images:
                result = mammoth.convert_to_html(docx_file, convert_image=mammoth.images.img_element(_convert_image_to_base64))
            else:
                result = mammoth.convert_to_html(docx_file)

            html = result.value

            markdown_text = markdownify(html)
            return markdown_text

        finally:
            if binary is None:
                docx_file.close()


class Pdf(PdfParser):
    def __init__(self):
        super().__init__()

    def __call__(self, filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, zoomin=3, callback=None, separate_tables_figures=False):
        start = timer()
        first_start = start
        callback(msg="OCR started")
        self.__images__(filename if binary is None else binary, zoomin, from_page, to_page, callback)
        callback(msg="OCR finished ({:.2f}s)".format(timer() - start))
        logging.info("OCR({}~{}): {:.2f}s".format(from_page, to_page, timer() - start))

        start = timer()
        self._layouts_rec(zoomin)
        callback(0.63, "Layout analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._table_transformer_job(zoomin)
        callback(0.65, "Table analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._text_merge(zoomin=zoomin)
        callback(0.67, "Text merged ({:.2f}s)".format(timer() - start))

        if separate_tables_figures:
            tbls, figures = self._extract_table_figure(True, zoomin, True, True, True)
            self._concat_downward()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls, figures
        else:
            tbls = self._extract_table_figure(True, zoomin, True, True)
            self._naive_vertical_merge()
            self._concat_downward()
            # self._final_reading_order_merge()
            # self._filter_forpages()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls


# Maximum number of HTTP redirects followed when fetching a remote image
# referenced by a markdown document (each hop is SSRF-validated).
MAX_IMAGE_REDIRECTS = 5


class Markdown(MarkdownParser):
    def md_to_html(self, sections):
        if not sections:
            return []
        if isinstance(sections, type("")):
            text = sections
        elif isinstance(sections[0], type("")):
            text = sections[0]
        else:
            return []

        from bs4 import BeautifulSoup

        html_content = markdown(text)
        soup = BeautifulSoup(html_content, "html.parser")
        return soup

    def get_hyperlink_urls(self, soup):
        if soup:
            return set([a.get("href") for a in soup.find_all("a") if a.get("href")])
        return []

    def extract_image_urls_with_lines(self, text):
        md_img_re = re.compile(r"!\[[^\]]*\]\(([^)\s]+)")
        html_img_re = re.compile(r'src=["\\\']([^"\\\'>\\s]+)', re.IGNORECASE)
        urls = []
        seen = set()
        lines = text.splitlines()
        for idx, line in enumerate(lines):
            for url in md_img_re.findall(line):
                if (url, idx) not in seen:
                    urls.append({"url": url, "line": idx})
                    seen.add((url, idx))
            for url in html_img_re.findall(line):
                if (url, idx) not in seen:
                    urls.append({"url": url, "line": idx})
                    seen.add((url, idx))

        # cross-line
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "html.parser")
            newline_offsets = [m.start() for m in re.finditer(r"\n", text)] + [len(text)]
            for img_tag in soup.find_all("img"):
                src = img_tag.get("src")
                if not src:
                    continue

                tag_str = str(img_tag)
                pos = text.find(tag_str)
                if pos == -1:
                    # fallback
                    pos = max(text.find(src), 0)
                line_no = 0
                for i, off in enumerate(newline_offsets):
                    if pos <= off:
                        line_no = i
                        break
                if (src, line_no) not in seen:
                    urls.append({"url": src, "line": line_no})
                    seen.add((src, line_no))
        except Exception as e:
            logging.error("Failed to extract image urls: {}".format(e))
            pass

        return urls

    def load_images_from_urls(self, urls, cache=None):
        import requests
        from pathlib import Path
        from urllib.parse import urljoin

        from common.ssrf_guard import assert_url_is_safe, pin_dns

        cache = cache or {}
        images = []
        for url in urls:
            if url in cache:
                if cache[url]:
                    images.append(cache[url])
                continue
            img_obj = None
            try:
                if url.startswith(("http://", "https://")):
                    # SSRF guard: image references come from the (untrusted) uploaded
                    # document, so validate and DNS-pin every hop before connecting.
                    # Otherwise a markdown image like ![x](http://169.254.169.254/...)
                    # would make the server fetch internal services / cloud metadata.
                    # Redirects are followed manually so each hop is re-validated,
                    # mirroring common/data_source/rss_connector.py.
                    current_hostname, current_ip = assert_url_is_safe(url)
                    current_url = url
                    response = None
                    try:
                        for _ in range(MAX_IMAGE_REDIRECTS + 1):
                            # Release the previous hop before opening the next: with
                            # stream=True the connection isn't returned to the pool
                            # until the body is read or the response is closed.
                            if response is not None:
                                response.close()
                            with pin_dns(current_hostname, current_ip):
                                response = requests.get(current_url, stream=True, timeout=30, allow_redirects=False)
                            if response.status_code not in (301, 302, 303, 307, 308):
                                break
                            location = response.headers.get("Location")
                            if not location:
                                break
                            current_url = urljoin(current_url, location)
                            current_hostname, current_ip = assert_url_is_safe(current_url)
                        else:
                            raise ValueError(f"Exceeded {MAX_IMAGE_REDIRECTS} redirects fetching {url!r}")
                        if response.status_code == 200 and response.headers.get("Content-Type", "").startswith("image/"):
                            img_obj = Image.open(BytesIO(response.content)).convert("RGB")
                    finally:
                        # Always release the final/streamed response, including the
                        # non-image and redirect-cap paths where the body is unread.
                        if response is not None:
                            response.close()
                else:
                    local_path = Path(url)
                    if local_path.exists():
                        img_obj = Image.open(url).convert("RGB")
                    else:
                        logging.warning(f"Local image file not found: {url}")
            except Exception as e:
                logging.error(f"Failed to download/open image from {url}: {e}")
            cache[url] = img_obj
            if img_obj:
                images.append(img_obj)
        return images, cache

    def __call__(self, filename, binary=None, separate_tables=True, delimiter=None, return_section_images=False):
        """Parse markdown into text sections and optional standalone table chunks."""
        if binary is not None:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(filename, "r") as f:
                txt = f.read()

        remainder, tables = self.extract_tables_and_remainder(f"{txt}\n", separate_tables=separate_tables)
        parsing_text = remainder
        extractor = MarkdownElementExtractor(parsing_text)
        image_refs = self.extract_image_urls_with_lines(parsing_text)
        element_sections = extractor.extract_elements(delimiter, include_meta=True)

        sections = []
        section_images = []
        image_cache = {}
        for element in element_sections:
            content = element["content"]
            start_line = element["start_line"]
            end_line = element["end_line"]
            urls_in_section = [ref["url"] for ref in image_refs if start_line <= ref["line"] <= end_line]
            imgs = []
            if urls_in_section:
                imgs, image_cache = self.load_images_from_urls(urls_in_section, image_cache)
            combined_image = None
            if imgs:
                combined_image = reduce(concat_img, imgs) if len(imgs) > 1 else imgs[0]
            sections.append((content, ""))
            section_images.append(combined_image)

        tbls = []
        if separate_tables:
            for table in tables:
                tbls.append(((None, markdown(table, extensions=["markdown.extensions.tables"])), ""))
        if return_section_images:
            return sections, tbls, section_images
        return sections, tbls


def load_from_xml_v2(baseURI, rels_item_xml):
    """
    Return |_SerializedRelationships| instance loaded with the
    relationships contained in *rels_item_xml*. Returns an empty
    collection if *rels_item_xml* is |None|.
    """
    srels = _SerializedRelationships()
    if rels_item_xml is not None:
        rels_elm = parse_xml(rels_item_xml)
        for rel_elm in rels_elm.Relationship_lst:
            if rel_elm.target_ref in ("../NULL", "NULL") or rel_elm.target_ref.startswith("#"):
                continue
            srels._srels.append(_SerializedRelationship(baseURI, rel_elm))
    return srels
# general 对应的chunk
# binary 是从 MinIO 读取的整个原文件的字节内容，不是指定页段的文件内容。
# from_page/to_page 与完整 binary 一起传给解析器；以默认 DeepDOC PDF 解析器为例，它打开完整 PDF 后，再通过 pdf.pages[from_page:to_page] 选择本 Task 要处理的页（结束页不包含）
def chunk(filename, binary=None, from_page=0, to_page=MAXIMUM_PAGE_NUMBER, lang="Chinese", callback=None, **kwargs):
    """按 general/naive 策略解析原文件，并返回尚未生成向量的 Chunk 字典列表。

    filename 决定文件类型；binary 是原文件内容。PDF 解析器会接收当前 Task 的
    [from_page, to_page) 范围，但具体是否裁页取决于所选实现。lang 用于分词和
    解析器语言；callback 由 Worker 传入以报告进度；kwargs 包含 parser_config、
    tenant_id、kb_id 等上下文。返回值只在内存中，尚未生成 Embedding。
    """
    # 收集文档中的超链接；链接内容解析出的 Chunk 最后并入当前结果。
    urls = set()
    url_res = []

    # 语言缺省为中文；parser_config 决定解析器、分隔符及目标 Chunk 大小。
    lang = lang or "Chinese"
    # 向下游分词/Chunk 构造函数传递是否为英文的标记。
    is_english = lang.lower() == "english"
    # 只有完全未传 parser_config 时使用这里的整套默认值；已传字典的缺失键在各分支单独取默认值。
    parser_config = kwargs.get("parser_config", {"chunk_token_num": 512, "delimiter": "\n!?。；！？", "layout_recognize": "DeepDOC", "analyze_hyperlink": True})

    # children_delimiter 是主 Chunk 生成后的二次拆分规则；先还原转义字符。
    child_deli = (parser_config.get("children_delimiter") or "").encode("utf-8").decode("unicode_escape").encode("latin1").decode("utf-8")
    # 反引号包裹的多字符分隔符按整体匹配，其余字符逐个组成正则分隔规则。
    cust_child_deli = re.findall(r"`([^`]+)`", child_deli)
    child_deli = "|".join(re.sub(r"`([^`]+)`", "", child_deli))
    if cust_child_deli:
        # 长串优先匹配，避免短串截断同一位置的长分隔符。
        cust_child_deli = sorted(set(cust_child_deli), key=lambda x: -len(x))
        cust_child_deli = "|".join(re.escape(t) for t in cust_child_deli if t)
        child_deli += cust_child_deli

    # Markdown 有独立的合并流程；表格/图片上下文窗口用于给媒体块补相邻正文。
    is_markdown = False
    table_context_size = max(0, int(parser_config.get("table_context_size", 0) or 0))
    image_context_size = max(0, int(parser_config.get("image_context_size", 0) or 0))

    # 基础字段会复制到每个 Chunk：原文件名、标题粗粒度/细粒度分词。
    doc = {"docnm_kwd": filename, "title_tks": rag_tokenizer.tokenize(re.sub(r"\.[a-zA-Z]+$", "", filename))}
    doc["title_sm_tks"] = rag_tokenizer.fine_grained_tokenize(doc["title_tks"])
    # res 保存当前文件解析出的 Chunk；pdf_parser 后续用于裁图和提取位置。
    res = []
    pdf_parser = None
    section_images = None

    # 只在顶层文件提取嵌入附件，子文件不再递归扫描，避免重复处理。
    is_root = kwargs.get("is_root", True)
    embed_res = []
    if is_root:
        embeds = []
        if binary is not None:
            # 从原始字节中提取嵌入文件，不是从 MinIO 再读取一次。
            embeds = extract_embed_file(binary)
        else:
            raise Exception("Embedding extraction from file path is not supported.")

        # 每个附件也走本方法，但标记 is_root=False；单个附件失败不终止主文件。
        for embed_filename, embed_bytes in embeds:
            try:
                sub_res = chunk(embed_filename, binary=embed_bytes, lang=lang, callback=callback, is_root=False, **kwargs) or []
                embed_res.extend(sub_res)
            except Exception as e:
                error_msg = f"Failed to chunk embed {embed_filename}: {e}"
                logging.error(error_msg)
                if callback:
                    callback(0.05, error_msg)
                continue

    # DOCX 走专用合并器，处理完后直接返回；此分支也未使用 from_page/to_page。
    if re.search(r"\.docx$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        # 配置允许时，抓取 DOCX 超链接指向的网页并作为额外内容解析。
        if parser_config.get("analyze_hyperlink", False) and is_root:
            urls = extract_links_from_docx(binary)
            for index, url in enumerate(urls):
                html_bytes, metadata = extract_html(url)
                if not html_bytes:
                    continue
                try:
                    # URL 后缀可识别时直接按其类型解析。
                    sub_url_res = chunk(url, html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
                except Exception as e:
                    logging.info(f"Failed to chunk url in registered file type {url}: {e}")
                    # 原 URL 的解析失败时，改用 HTML 文件名重试。
                    sub_url_res = chunk(f"{index}.html", html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
                url_res.extend(sub_url_res)

        # 修正 python-docx 读取含无效 word/NULL 关系的文档时可能抛出的异常。
        _SerializedRelationships.load_from_xml = load_from_xml_v2

        # 解析 DOCX 的正文、图片和表格，并规范化可能存在的 RTL 表现形式字符。
        sections = Docx()(filename, binary)
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        # 依据目标 token 数和分隔符合并 DOCX 块；images 标识需增强的图片块下标。
        chunks, images = naive_merge_docx(sections, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), table_context_size, image_context_size)

        # 有视觉模型等配置时，为图片块补充图像理解内容。
        vision_figure_parser_docx_wrapper_naive(
            chunks=chunks,
            idx_lst=images,
            callback=callback,
            lang=lang,
            **kwargs,
        )

        callback(0.8, "Finish parsing.")
        st = timer()

        # 将 DOCX 块转成检索字段，再并入附件和链接块；本分支不会应用 overlapped_percent。
        res.extend(doc_tokenize_chunks_with_images(chunks, doc, is_english, child_delimiters_pattern=child_deli, language=lang))
        logging.info("naive_merge({}): {}".format(filename, timer() - st))
        res.extend(embed_res)
        res.extend(url_res)
        return res

    # PDF 先选择版面解析器，再把正文与表格分别转换成 Chunk。
    elif re.search(r"\.pdf$", filename, re.IGNORECASE):
        # layout_recognize 可以是内置解析器名称，也可以是租户配置的模型 ID。
        layout_recognize_raw = parser_config.get("layout_recognize", "DeepDOC")
        tenant_id = kwargs.get("tenant_id")
        if tenant_id and isinstance(layout_recognize_raw, str):
            try:
                # 将模型 ID 还原为带供应商信息的组合名称，供分发器识别。
                layout_recognize_raw = get_composite_model_name_by_id(layout_recognize_raw)
            except LookupError:
                # 普通解析器名称或无法解析的 ID 原样交给分发器处理。
                pass
        # 分发器内部负责规范化名称；这里不重复规范化，以免丢失模型名称和供应商后缀。
        opendataloader_llm_name = kwargs.pop("opendataloader_llm_name", None)
        # 得到实际 PDF 解析函数、解析器名称及其可能依赖的模型配置。
        parser, name, layout_recognizer, opendataloader_llm_name, parser_model_name = _dispatch_pdf_parser(
            parser_config,
            opendataloader_llm_name,
            layout_recognize_override=layout_recognize_raw,
        )

        # 可选：提取 PDF 中的超链接，网页内容会在正文切块后另行解析。
        if parser_config.get("analyze_hyperlink", False) and is_root:
            urls = extract_links_from_pdf(binary)
        callback(0.1, "Start to parse.")
        if name == "mineru":
            # general 策略下的 MinerU 使用 naive 解析模式。
            kwargs["parse_method"] = "naive"

        # 真正执行 PDF 解析：sections 为正文及位置，tables 为表格/图片结果。
        # from_page/to_page 传给选中的解析器；具体是否裁页由该解析器实现决定。
        # 多种 *_llm_name 参数共用已解析的模型名称，由选中的解析器读取对应项。
        sections, tables, pdf_parser = parser(
            filename=filename,
            binary=binary,
            from_page=from_page,
            to_page=to_page,
            lang=lang,
            callback=callback,
            layout_recognizer=layout_recognizer,
            mineru_llm_name=parser_model_name,
            paddleocr_llm_name=parser_model_name,
            opendataloader_llm_name=opendataloader_llm_name,
            somark_llm_name=parser_model_name,
            mistral_ocr_llm_name=parser_model_name,
            **kwargs,
        )
        # 统一处理阿拉伯语等 RTL 字形，避免后续分词保留表现形式字符。
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        # 当前页段没有可解析的正文或表格，则直接返回空列表（附件结果也不追加）。
        if not sections and not tables:
            return []

        # 可选：为 PDF 表格/图片附上邻近正文；实际窗口长度只取 image_context_size。
        # 仅 table_context_size 非零而 image_context_size 为零时，调用不会增加上下文。
        if table_context_size or image_context_size:
            tables = append_context2table_image4pdf(
                sections,
                tables,
                image_context_size,
                section_page_offset=from_page if name == "mineru" else 0,
            )

        # 对特定外部解析器，非正数 token 目标表示不再把输出段合并成较大正文块。
        # 这里直接修改传入的 parser_config 字典，后面的合并逻辑会读取新值。
        if name in ["tcadp", "docling", "mineru", "paddleocr", "opendataloader", "somark", "mistral ocr"]:
            if int(parser_config.get("chunk_token_num", 0)) <= 0:
                parser_config["chunk_token_num"] = 0

        # 表格/图片先独立转成检索 Chunk；正文 sections 稍后进入通用合并流程。
        res = tokenize_table(tables, doc, is_english, language=lang)
        callback(0.8, "Finish parsing.")

    # CSV/XLS/XLSX 在 general 策略下使用通用表格读取，不等于专门的 table 解析策略。
    elif re.search(r"\.(csv|xlsx?)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")

        # 配置选择 TCADP 时走外部表格解析器，否则使用默认 ExcelParser。
        layout_recognizer = parser_config.get("layout_recognize", "DeepDOC")
        if layout_recognizer == "TCADP Parser":
            table_result_type = parser_config.get("table_result_type", "1")
            markdown_image_response_type = parser_config.get("markdown_image_response_type", "1")
            tcadp_parser = TCADPParser(table_result_type=table_result_type, markdown_image_response_type=markdown_image_response_type)
            # 依赖不可用时报告失败并返回空结果。
            if not tcadp_parser.check_installation():
                callback(-1, "TCADP parser not available. Please check Tencent Cloud API configuration.")
                return res

            # TCADP 接口通过 file_type 区分 Excel 与 CSV。
            file_type = "XLSX" if re.search(r"\.xlsx?$", filename, re.IGNORECASE) else "CSV"

            # 解析出正文和表格；表格立即转成 Chunk，正文留待后续合并。
            sections, tables = tcadp_parser.parse_pdf(filepath=filename, binary=binary, callback=callback, output_dir=os.environ.get("TCADP_OUTPUT_DIR", ""), file_type=file_type)
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)
            parser_config["chunk_token_num"] = 0
            res = tokenize_table(tables, doc, is_english, language=lang)
            callback(0.8, "Finish parsing.")
        else:
            # 默认读取工作表内容；html4excel 将表格组织成 HTML 片段。
            excel_parser = ExcelParser()
            if parser_config.get("html4excel"):
                sections = [(_, "") for _ in excel_parser.html(binary, 12) if _]
                # HTML 表格片段保持独立，不再合并多个片段。
                parser_config["chunk_token_num"] = 0
            else:
                sections = [(_, "") for _ in excel_parser(binary) if _]
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)

    # 纯文本与代码文件先由 TxtParser 拆出文本段，再进入公共合并流程。
    elif re.search(r"\.(txt|py|js|java|c|cpp|h|php|go|ts|sh|cs|kt|sql)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        sections = TxtParser()(filename, binary, parser_config.get("chunk_token_num", 128), parser_config.get("delimiter", "\n!?;。；！？"))
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        logging.info("TxtParser produced %d sections for %s", len(sections), filename)
        callback(0.8, "Finish parsing.")

    # Markdown 保留段落对应的图片，后面使用专用的按 section 合并流程。
    elif re.search(r"\.(md|markdown|mdx)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        markdown_parser = Markdown(int(parser_config.get("chunk_token_num", 128)))
        # sections 是正文段，tables 是独立媒体结果，section_images 与正文段下标对齐。
        sections, tables, section_images = markdown_parser(
            filename,
            binary,
            separate_tables=False,
            delimiter=parser_config.get("delimiter", "\n!?;。；！？"),
            return_section_images=True,
        )
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)

        # 标记走 Markdown 专用合并器，而非下方普通 naive_merge。
        is_markdown = True

        # 尝试获取租户默认视觉模型；缺失时继续解析文本，不中断导入。
        try:
            vision_model_config = get_tenant_default_model_by_type(kwargs["tenant_id"], LLMType.VISION)
            vision_model = LLMBundle(kwargs["tenant_id"], vision_model_config, lang=lang)
            callback(0.2, "Visual model detected. Attempting to enhance figure extraction...")
        except Exception as e:
            logging.warning(f"Failed to detect figure extraction: {e}")
            vision_model = None

        if vision_model:
            # 对每个带图片的正文段，用视觉模型生成图像描述并追加到段落文本。
            for idx, (section_text, _) in enumerate(sections):
                images = []
                if section_images and len(section_images) > idx and section_images[idx] is not None:
                    images.append(section_images[idx])

                if images and len(images) > 0:
                    # 同一段的多张图片先拼接，再作为一次视觉解析的输入。
                    combined_image = reduce(concat_img, images) if len(images) > 1 else images[0]
                    if section_images:
                        section_images[idx] = combined_image
                    else:
                        section_images = [None] * len(sections)
                        section_images[idx] = combined_image
                    markdown_vision_parser = VisionFigureParser(
                        vision_model=vision_model,
                        figures_data=[((combined_image, ["markdown image"]), [(0, 0, 0, 0, 0)])],
                        lang=lang,
                        **kwargs,
                    )
                    boosted_figures = markdown_vision_parser(callback=callback)
                    # 图像描述并入正文，后续同正文一起参与切块与检索分词。
                    sections[idx] = (section_text + "\n\n" + "\n\n".join([fig[0][1] for fig in boosted_figures]), sections[idx][1])

        else:
            logging.warning("No visual model detected. Skipping figure parsing enhancement.")

        # 仅顶层调用提取 Markdown 中的链接；是否抓取还要看 analyze_hyperlink。
        if parser_config.get("hyperlink_urls", False) and is_root:
            for idx, (section_text, _) in enumerate(sections):
                soup = markdown_parser.md_to_html(section_text)
                hyperlink_urls = markdown_parser.get_hyperlink_urls(soup)
                urls.update(hyperlink_urls)
        # 独立表格/图片结果先进入 res；Markdown 正文随后单独合并。
        res = tokenize_table(tables, doc, is_english, language=lang)
        callback(0.8, "Finish parsing.")

    # HTML/EPUB/JSON 系列先解出文本段，随后统一使用 naive_merge。
    elif re.search(r"\.(htm|html)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        sections = HtmlParser()(filename, binary, chunk_token_num)
        sections = [(_, "") for _ in sections if _]
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        callback(0.8, "Finish parsing.")

    # EPUB 解析结果也包装为 (文本, 位置) 形式；此处位置为空。
    elif re.search(r"\.epub$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        sections = EpubParser()(filename, binary, chunk_token_num)
        sections = [(_, "") for _ in sections if _]
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        callback(0.8, "Finish parsing.")

    # JSON/JSONL/LDJSON 按专用解析器提取文本，而非直接以原始 JSON 字符串切块。
    elif re.search(r"\.(json|jsonl|ldjson)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        sections = JsonParser(chunk_token_num)(binary)
        sections = [(_, "") for _ in sections if _]
        sections = _normalize_section_text_for_rtl_presentation_forms(sections)
        callback(0.8, "Finish parsing.")

    # 旧版 .doc 依赖 tika；不可用或未提取到正文时返回空列表。
    elif re.search(r"\.doc$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")

        try:
            from tika import parser as tika_parser
        except Exception as e:
            callback(0.8, f"tika not available: {e}. Unsupported .doc parsing.")
            logging.warning(f"tika not available: {e}. Unsupported .doc parsing for {filename}.")
            return []

        # Tika 从内存中的文件字节提取文本，按换行转成 sections。
        binary = BytesIO(binary)
        doc_parsed = tika_parser.from_buffer(binary)
        if doc_parsed.get("content", None) is not None:
            sections = doc_parsed["content"].split("\n")
            sections = [(_, "") for _ in sections if _]
            sections = _normalize_section_text_for_rtl_presentation_forms(sections)
            callback(0.8, "Finish parsing.")
        else:
            error_msg = f"tika.parser got empty content from {filename}."
            callback(0.8, error_msg)
            logging.warning(error_msg)
            return []
    else:
        # general 解析器不支持当前后缀；此处不负责转换文件类型。
        raise NotImplementedError("file type not supported yet(pdf, xlsx, doc, docx, txt supported)")

    # 文件类型解析完毕；以下将 sections 合并成 Chunk，并记录合并耗时。
    st = timer()
    # 重叠比例归一化到 0～90；它只作用于本次调用生成的相邻 Chunk。
    overlapped_percent = normalize_overlapped_percent(parser_config.get("overlapped_percent", 0))

    # Markdown 的段落与图片需要保持对齐，因此不走普通 naive_merge。
    if is_markdown:
        merged_chunks = []
        merged_images = []
        # chunk_token_num 是合并的目标值，不保证最终 Chunk 严格小于它。
        chunk_limit = max(0, int(parser_config.get("chunk_token_num", 128)))

        # 累积当前 Chunk 的正文、估计 token 数及合并后的图片。
        current_text = ""
        current_tokens = 0
        current_image = None

        for idx, sec in enumerate(sections):
            # Markdown 解析器的每个 section 通常是 (文本, 位置信息)。
            text = sec[0] if isinstance(sec, tuple) else sec
            sec_tokens = num_tokens_from_string(text)
            sec_image = section_images[idx] if section_images and idx < len(section_images) else None

            # 目标大小将被超过时，先封存旧 Chunk；短标题例外，尽量与后文同块。
            if current_text and not _is_short_header(current_text) and current_tokens + sec_tokens > chunk_limit:
                merged_chunks.append(current_text)
                merged_images.append(current_image)
                # 从旧块尾部按字符比例取重叠文本，作为下一块的开头。
                overlap_part = ""
                if overlapped_percent > 0:
                    overlap_len = int(len(current_text) * overlapped_percent / 100)
                    if overlap_len > 0:
                        overlap_part = current_text[-overlap_len:]
                # 有重叠时保留原图片引用；无重叠时从空块开始。
                current_text = overlap_part
                current_tokens = num_tokens_from_string(current_text)
                current_image = current_image if overlap_part else None

            # 将当前 section 拼到正在构造的 Chunk。
            if current_text:
                current_text += "\n" + text
            else:
                current_text = text
            current_tokens += sec_tokens

            # 当前 section 带图时，把它与本 Chunk 已收集的图片拼接。
            if sec_image:
                current_image = concat_img(current_image, sec_image) if current_image else sec_image

        # 循环结束后，最后一个尚未封存的 Chunk 也要加入结果。
        if current_text:
            merged_chunks.append(current_text)
            merged_images.append(current_image)

        chunks = merged_chunks
        has_images = merged_images and any(img is not None for img in merged_images)

        # 转换为含正文、分词等字段的字典；有图片时同时附上图片对象。
        if has_images:
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, merged_images, child_delimiters_pattern=child_deli, language=lang))
        else:
            res.extend(tokenize_chunks(chunks, doc, is_english, pdf_parser, child_delimiters_pattern=child_deli, language=lang))
    else:
        # 对非 Markdown 格式，空图片列表视为纯文本，不进入图文合并分支。
        if section_images:
            if all(image is None for image in section_images):
                section_images = None

        if section_images:
            # 图文路径：同步合并文本和图片，保持两者下标一一对应。
            chunks, images = naive_merge_with_images(sections, section_images, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), overlapped_percent)
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images, child_delimiters_pattern=child_deli, language=lang))
        else:
            # 纯文本路径：按 delimiter 拆段、按 token 软目标合并。
            # 普通路径可加入重叠；反引号自定义分隔符则每段独立成块，不应用重叠。
            chunks = naive_merge(sections, int(parser_config.get("chunk_token_num", 128)), parser_config.get("delimiter", "\n!?。；！？"), overlapped_percent)

            # 每个块生成正文和检索分词；PDF 还可从版面位置裁图、记录页码。
            res.extend(tokenize_chunks(chunks, doc, is_english, pdf_parser, child_delimiters_pattern=child_deli, language=lang))

    # 解析链接内容并追加为额外 Chunk；仅顶层且启用 analyze_hyperlink 时执行。
    if urls and parser_config.get("analyze_hyperlink", False) and is_root:
        for index, url in enumerate(urls):
            html_bytes, metadata = extract_html(url)
            if not html_bytes:
                continue
            try:
                # 先按 URL 的文件名/后缀选择子解析器。
                sub_url_res = chunk(url, html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
            except Exception as e:
                logging.info(f"Failed to chunk url in registered file type {url}: {e}")
                # 原 URL 的解析失败时，按 HTML 文档名重试。
                sub_url_res = chunk(f"{index}.html", html_bytes, callback=callback, lang=lang, is_root=False, **kwargs)
            url_res.extend(sub_url_res)

    logging.info("naive_merge({}): {}".format(filename, timer() - st))

    # 附件和超链接不是当前正文 sections 的一部分，在末尾合并进返回列表。
    if embed_res:
        res.extend(embed_res)
    if url_res:
        res.extend(url_res)

    # PDF 目录临时放在首个 Chunk 的 __outline__；后续任务流程会提取并单独保存。
    if res and pdf_parser and getattr(pdf_parser, "outlines", None):
        res[0]["__outline__"] = [{"title": title, "depth": depth} for title, depth, *_ in pdf_parser.outlines]

    # 这里仅返回内存字典；稳定 Chunk ID、Embedding 和检索索引写入都在调用方。
    return res


if __name__ == "__main__":
    import sys

    def dummy(prog=None, msg=""):
        pass

    chunk(sys.argv[1], from_page=0, to_page=10, callback=dummy)
